
import os
import subprocess
import fitz  # PyMuPDF
import docx
import csv
import re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import gspread
from google.oauth2.service_account import Credentials

# 📁 Setup paths
base_dir = "/Users/aneel/Desktop/ResumeMatcherAI"
resume_folder = os.path.join(base_dir, "resumes")
jd_folder = os.path.join(base_dir, "jds")
output_folder = os.path.join(base_dir, "output")
prompt_path = os.path.join(base_dir, "prompt.txt")
csv_output = os.path.join(output_folder, "scores.csv")
excel_output = os.path.join(output_folder, "scores.xlsx")

# 📄 Extract resume text
def extract_text(file_path):
    if file_path.endswith(".txt"):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    elif file_path.endswith(".pdf"):
        doc = fitz.open(file_path)
        return "".join([page.get_text() for page in doc])
    elif file_path.endswith(".docx"):
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    return ""

# 📥 Load JD and prompt
jd_files = [f for f in os.listdir(jd_folder) if f.endswith(".txt")]
if not jd_files:
    print("❌ No JD file found.")
    exit()
jd_text = extract_text(os.path.join(jd_folder, jd_files[0]))

with open(prompt_path, "r", encoding="utf-8") as f:
    prompt_template = f.read()

os.makedirs(output_folder, exist_ok=True)
csv_rows = [["Filename", "Score", "Reason"]]

# 🔁 Score resumes
resume_files = [f for f in os.listdir(resume_folder) if f.endswith((".pdf", ".docx", ".txt"))]
if not resume_files:
    print("❌ No resumes found.")
    exit()

for resume_file in resume_files:
    resume_path = os.path.join(resume_folder, resume_file)
    resume_text = extract_text(resume_path)

    final_prompt = prompt_template.replace("{{resume_text}}", resume_text).replace("{{jd_text}}", jd_text)
    print(f"\n🔍 Scoring: {resume_file}")

    result = subprocess.run(
        ["ollama", "run", "mistral"],  # or "llama3"
        input=final_prompt.encode("utf-8"),
        stdout=subprocess.PIPE
    )

    decoded = result.stdout.decode("utf-8")
    print(f"✅ Done: {resume_file}\n→ AI Says:\n{decoded.strip()}")

    with open(os.path.join(output_folder, f"result_{resume_file}.txt"), "w", encoding="utf-8") as f:
        f.write(decoded)

    score_match = re.search(r"(?i)score:\s*(\d+)", decoded)
    reason_match = re.search(r"(?i)reason:\s*(.+)", decoded)

    score_line = score_match.group(1) if score_match else ""
    reason_line = reason_match.group(1).strip() if reason_match else ""
    csv_rows.append([resume_file, score_line, reason_line])

# 💾 Save CSV
with open(csv_output, "w", newline='', encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerows(csv_rows)

# 💾 Save Excel
wb = Workbook()
ws = wb.active
ws.title = "Resume Match Scores"
ws.append(csv_rows[0])
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center")
for row in csv_rows[1:]:
    ws.append(row)
wb.save(excel_output)

print(f"\n📁 CSV saved to: {csv_output}")
print(f"📊 Excel saved to: {excel_output}")

# 📤 Upload to Google Sheets
def upload_to_gsheets(data, sheet_name="Resume Scores"):
    creds_path = os.path.join(base_dir, "credentials.json")
    scope = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]
    creds = Credentials.from_service_account_file(creds_path, scopes=scope)
    client = gspread.authorize(creds)

    try:
        sheet = client.open(sheet_name)
    except gspread.SpreadsheetNotFound:
        sheet = client.create(sheet_name)
        sheet.share("aneelv75@gmail.com", perm_type="user", role="writer")

    worksheet = sheet.sheet1
    worksheet.clear()
    worksheet.update(values=data, range_name="A1")  # ✅ FIXED: no indentation issue

    print(f"✅ Google Sheet updated: https://docs.google.com/spreadsheets/d/{sheet.id}")

# ⬆ Push to Google Sheets
upload_to_gsheets(csv_rows)

def upload_to_gsheets(data):
    creds_path = os.path.join(base_dir, "credentials.json")
    scope = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]
    creds = Credentials.from_service_account_file(creds_path, scopes=scope)
    client = gspread.authorize(creds)

    # 📌 OPEN SHEET BY KEY
    sheet = client.open_by_key("1rMBAkIPBRuEMk91-HyQ_Adgu3ZltRc_cgHUIuJLWdIU")

    try:
        worksheet = sheet.sheet1
        worksheet.clear()
        worksheet.update(values=data, range_name="A1")

        print("✅ Google Sheet updated:")
        print("🔗 https://docs.google.com/spreadsheets/d/1rMBAkIPBRuEMk91-HyQ_Adgu3ZltRc_cgHUIuJLWdIU")
    except Exception as e:
        print("❌ Error updating sheet:", e)


